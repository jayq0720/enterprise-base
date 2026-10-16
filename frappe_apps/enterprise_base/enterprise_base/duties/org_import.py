"""组织任职导入：服务器全量预检，原生文档写入，整批事务回滚。"""

import datetime as dt
import io
import zipfile
from html import escape

import frappe
from openpyxl import load_workbook

COLUMNS = {
    "部门": ["名称", "上级部门"],
    "岗位": ["名称", "岗位职责", "可见的知识空间"],
    "员工任职": ["工号", "姓名", "企业微信 userid", "部门", "岗位", "直属上级工号", "入职日期"],
}


def read_rows(content):
    errors: list[dict] = []
    rows: dict[str, list[tuple[int, list[str]]]] = {name: [] for name in COLUMNS}

    def error(sheet, row, column, message):
        errors.append(dict(sheet=sheet, row=row, column=column, message=message))

    try:
        if len(content) > 2_000_000:
            raise ValueError
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 100
                or sum(i.file_size for i in archive.infolist()) > 20_000_000
            ):
                raise ValueError
            if any("vbaProject" in n or "externalLinks" in n for n in archive.namelist()):
                raise ValueError
        book = load_workbook(io.BytesIO(content), read_only=False, data_only=False)
        if book.sheetnames != [*COLUMNS, "说明"]:
            raise ValueError
        if [c.value for c in book["说明"][1]] != ["工作表", "列名或事项", "必填", "规则", "示例"]:
            raise ValueError
        for sheet, columns in COLUMNS.items():
            tab = book[sheet]
            if tab.max_row > 1001 or tab.max_column != len(columns) or tab.merged_cells:
                error(sheet, 1, "表结构", "最多 1000 行，不得增加列或合并单元格。")
                continue
            if [c.value for c in tab[1]] != columns:
                error(sheet, 1, "列名", "列名必须与模板完全一致；禁止敏感字段列。")
                continue
            seen = set()
            for number, cells in enumerate(tab.iter_rows(min_row=2), 2):
                values = [str(c.value).strip() if c.value is not None else "" for c in cells]
                if not any(values):
                    continue
                for label, cell in zip(columns, cells, strict=True):
                    if cell.data_type == "f" or (
                        cell.value is not None and len(str(cell.value)) > 10000
                    ):
                        error(sheet, number, label, "不允许公式或超过 10000 字的单元格。")
                required = (
                    (0,) if sheet == "部门" else ((0, 1) if sheet == "岗位" else (0, 1, 2, 3, 4, 6))
                )
                for col in required:
                    if not values[col]:
                        error(sheet, number, columns[col], "不能为空。")
                if values[0] in seen:
                    error(sheet, number, columns[0], "文件内重复。")
                seen.add(values[0])
                if sheet != "员工任职" and "," in values[0]:
                    error(sheet, number, columns[0], "不能含英文逗号。")
                if sheet == "员工任职":
                    for col in (0, 2, 5):
                        if cells[col].value is not None and not isinstance(cells[col].value, str):
                            error(sheet, number, columns[col], "必须为文本，避免丢失前导零。")
                    try:
                        value = cells[6].value
                        values[6] = (
                            value.date()
                            if isinstance(value, dt.datetime)
                            else dt.date.fromisoformat(str(value))
                        ).isoformat()
                    except (TypeError, ValueError):
                        error(sheet, number, columns[6], "必须为有效日期 yyyy-mm-dd。")
                rows[sheet].append((number, values))
        book.close()
    except Exception:
        error("工作簿", 1, "文件", "文件必须为四表 xlsx，且未损坏、不含宏或外部工作簿链接。")
    return rows, errors


def prepare(company, rows, errors):
    def error(sheet, number, col, message):
        errors.append(dict(sheet=sheet, row=number, column=COLUMNS[sheet][col], message=message))

    departments = frappe.get_all(
        "Department",
        filters={"company": company},
        fields=["name", "department_name", "parent_department", "is_group"],
    )
    aliases = {d.department_name: d.name for d in departments}
    aliases.update({d.name: d.name for d in departments})
    for _, row in rows["部门"]:
        aliases[row[0]] = aliases.get(row[0], row[0])
    employees = frappe.get_all(
        "Employee",
        filters={"company": company},
        fields=["name", "employee_number", "employee_name", "wecom_userid", "reports_to"],
    )
    numbers: dict[str, list] = {}
    for employee in employees:
        numbers.setdefault(employee.employee_number, []).append(employee)
    employee_graph = {e.name: e.reports_to for e in employees}
    department_graph = {d.name: d.parent_department for d in departments}
    managers = frappe.session.user == "Administrator" or "HR Manager" in frappe.get_roles()
    uids, selected, used_departments = set(), {}, set()
    for number, row in rows["部门"]:
        if aliases[row[0]] in used_departments:
            error("部门", number, 0, "多个名称指向同一部门。")
        used_departments.add(aliases[row[0]])
        department_graph[aliases[row[0]]] = aliases.get(row[1]) if row[1] else None
        if row[1] and row[1] not in aliases:
            error("部门", number, 1, "上级部门不存在或不属于所选公司。")
        if any(d.name == aliases.get(row[1]) and not d.is_group for d in departments) and row[
            1
        ] not in {r[0] for _, r in rows["部门"]}:
            error("部门", number, 1, "上级部门须先设为群组，或在本文件中一并维护。")
    for number, row in rows["岗位"]:
        spaces = [s.strip() for s in row[2].split(",") if s.strip()]
        if (
            (row[2] and any(not s.strip() for s in row[2].split(",")))
            or len(spaces) != len(set(spaces))
            or any(not frappe.db.exists("EB Knowledge Space", s) for s in spaces)
        ):
            error("岗位", number, 2, "知识空间不存在或重复。")
        old = (
            frappe.get_doc("Designation", row[0])
            if frappe.db.exists("Designation", row[0])
            else None
        )
        if not managers and set(spaces) != {
            s.knowledge_space for s in (old.eb_knowledge_spaces if old else [])
        }:
            error("岗位", number, 2, "可见范围只能由 HR Manager 修改。")
    jobs = {row[0] for _, row in rows["岗位"]}
    for number, row in rows["员工任职"]:
        matches = numbers.get(row[0], [])
        if len(matches) != 1:
            error("员工任职", number, 0, "工号不存在或不唯一；请 HR 先在后台建档。")
            continue
        employee = matches[0]
        if "System Manager" not in frappe.get_roles() and not frappe.has_permission(
            "Employee", "write", doc=frappe.get_doc("Employee", employee.name)
        ):
            error("员工任职", number, 0, "无权修改该员工档案。")
        selected[row[0]] = employee.name
        if employee.employee_name != row[1]:
            error("员工任职", number, 1, "姓名与工号对应档案不符。")
        if row[2] in uids or frappe.db.exists(
            "Employee", {"wecom_userid": row[2], "name": ["!=", employee.name]}
        ):
            error("员工任职", number, 2, "企业微信 userid 已由其他员工使用。")
        uids.add(row[2])
        if not managers and (employee.wecom_userid or "") != row[2]:
            error("员工任职", number, 2, "身份关联只能由 HR Manager 修改。")
        if row[3] not in aliases:
            error("员工任职", number, 3, "部门不存在或不属于所选公司。")
        if row[4] not in jobs and not frappe.db.exists("Designation", row[4]):
            error("员工任职", number, 4, "岗位不存在。")
        boss = numbers.get(row[5], [])
        if row[5] and len(boss) != 1:
            error("员工任职", number, 5, "上级工号不存在或不唯一。")
        employee_graph[employee.name] = boss[0].name if len(boss) == 1 else None
    for sheet, graph, col in (("部门", department_graph, 1), ("员工任职", employee_graph, 5)):
        for number, row in rows[sheet]:
            node, visited = (aliases[row[0]] if sheet == "部门" else selected.get(row[0])), set()
            while node in graph and node not in visited:
                visited.add(node)
                node = graph[node]
            if node in visited:
                error(sheet, number, col, "上级关系不能指向自身或形成循环。")
    return aliases, numbers


@frappe.whitelist(methods=["POST"])
def import_workbook(file_url, company, save=False):
    frappe.only_for(["HR Manager", "System Manager"])
    frappe.get_doc("Company", company).check_permission("read")
    file = frappe.get_doc("File", {"file_url": file_url})
    file.check_permission("read")
    if (
        not file.is_private
        or file.file_size > 2_000_000
        or not file.file_name.lower().endswith(".xlsx")
    ):
        frappe.throw("请上传私有 xlsx 文件。")

    def persist(doc):
        if "System Manager" not in frappe.get_roles():
            doc.check_permission("create" if doc.is_new() else "write")
        doc.save(ignore_permissions=True)  # 已验角色、公司及字段白名单；文档校验仍执行。

    # 预检和保存都重新读文件、核对数据库，不信任前端提交的预检结果。
    with frappe.cache.lock(frappe.cache.make_key("eb-org-import"), timeout=300, blocking_timeout=3):
        rows, errors = read_rows(file.get_content())
        if errors:
            return {"errors": errors}
        aliases, numbers = prepare(company, rows, errors)
        if errors or not frappe.utils.cint(save):
            return {"errors": errors, "counts": {k: len(v) for k, v in rows.items()}}
        frappe.db.savepoint("eb_org_import")
        try:
            location = ("部门", 0)
            pending = list(rows["部门"])
            children = {aliases[r[1]] for _, r in pending if r[1]}
            while pending:
                for item in list(pending):
                    number, row = item
                    parent = aliases.get(row[1])
                    if parent and parent in {aliases[r[0]] for _, r in pending}:
                        continue
                    location = ("部门", number)
                    name = aliases[row[0]]
                    doc = (
                        frappe.get_doc("Department", name)
                        if frappe.db.exists("Department", {"name": name, "company": company})
                        else frappe.new_doc("Department")
                    )
                    doc.update(
                        {
                            "company": company,
                            "department_name": doc.department_name or row[0],
                            "parent_department": parent,
                        }
                    )
                    doc.is_group = int(bool(doc.is_group) or name in children)
                    persist(doc)
                    aliases[row[0]] = doc.name
                    pending.remove(item)
            for number, row in rows["岗位"]:
                location = ("岗位", number)
                doc = (
                    frappe.get_doc("Designation", row[0])
                    if frappe.db.exists("Designation", row[0])
                    else frappe.new_doc("Designation")
                )
                doc.designation_name = row[0]
                doc.eb_responsibilities = "".join(
                    f"<p>{escape(line)}</p>" for line in row[1].splitlines() if line.strip()
                )
                doc.set(
                    "eb_knowledge_spaces",
                    [{"knowledge_space": s.strip()} for s in row[2].split(",") if s.strip()],
                )
                persist(doc)
            for number, row in rows["员工任职"]:
                location = ("员工任职", number)
                doc = frappe.get_doc("Employee", numbers[row[0]][0].name)
                doc.update(
                    {
                        "wecom_userid": row[2],
                        "department": aliases[row[3]],
                        "designation": row[4],
                        "reports_to": numbers[row[5]][0].name if row[5] else None,
                        "date_of_joining": row[6],
                    }
                )
                persist(doc)
        except Exception:
            frappe.db.rollback(save_point="eb_org_import")
            # 原生校验可能带员工字段值；只返回位置，不暴露原始异常内容。
            return {
                "errors": [
                    dict(
                        sheet=location[0],
                        row=location[1],
                        column="保存",
                        message="原生文档校验失败，已整批回滚；请核对日期、部门层级及档案状态。",
                    )
                ]
            }
    return {"errors": [], "saved": True, "counts": {k: len(v) for k, v in rows.items()}}
