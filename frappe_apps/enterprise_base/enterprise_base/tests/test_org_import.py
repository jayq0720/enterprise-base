"""真实 Frappe 站点：全量拒绝、事务回滚、幂等及字段边界。"""

import io
import json
from pathlib import Path
from unittest.mock import patch

import frappe
from frappe.tests.utils import FrappeTestCase
from openpyxl import Workbook

from enterprise_base.duties.org_import import COLUMNS, import_workbook, read_rows


class TestOrgImport(FrappeTestCase):
    def setUp(self):
        frappe.set_user("Administrator")
        data = json.loads(Path(__file__).with_name("p2_data.json").read_text())
        for row in data["records"]:
            if row["doctype"] != "Version" and not frappe.db.exists(row["doctype"], row["name"]):
                frappe.get_doc(row).insert()
        self.suffix = frappe.generate_hash(length=8)
        employee = data["employee"]
        employee.update(first_name="Import", user_id=None, leave_approver=None)
        employee.update(employee_number=self.suffix, wecom_userid=self.suffix)
        self.emp = frappe.get_doc(employee).insert()
        self.rows = json.loads(
            Path(__file__).with_name("org_import_data.json").read_text().replace("$id", self.suffix)
        )

    def content(self, extra=None):
        book, output = Workbook(), io.BytesIO()
        book.remove(book.active)
        for (sheet, columns), rows in zip(COLUMNS.items(), self.rows, strict=True):
            tab = book.create_sheet(sheet)
            if extra and sheet == "员工任职":
                columns = [*columns, extra]
            for row in [columns, *rows]:
                tab.append(row)
        book.create_sheet("说明").append(["工作表", "列名或事项", "必填", "规则", "示例"])
        book.save(output)
        return output.getvalue()

    def run_import(self, save=True):
        filename = self.suffix + ".xlsx"
        file = frappe.get_doc(
            dict(doctype="File", file_name=filename, is_private=1, content=self.content())
        ).insert()
        return import_workbook(file.file_url, self.emp.company, save)

    def test_save_and_repeat(self):
        self.assertFalse(self.run_import(False)["errors"])
        self.assertFalse(frappe.db.exists("Designation", self.rows[1][0][0]))
        for _ in range(2):
            self.assertTrue(self.run_import()["saved"])
        self.emp.reload()
        self.assertEqual(self.emp.designation, self.rows[1][0][0])
        self.assertEqual(frappe.db.count("Designation", {"name": self.emp.designation}), 1)
        self.assertEqual(
            frappe.db.get_value("Designation", self.emp.designation, "eb_responsibilities"),
            "<p>第一项</p><p>第二项</p>",
        )

    def test_invalid_rows_leave_database_unchanged(self):
        self.rows[0][0][1] = self.rows[0][0][0]
        self.rows[1][0][2] = "missing-space"
        self.rows[2][0][0] = "missing-employee"
        result = self.run_import()
        self.assertEqual(
            {(e["sheet"], e["row"], e["column"]) for e in result["errors"]},
            {("部门", 2, "上级部门"), ("岗位", 2, "可见的知识空间"), ("员工任职", 2, "工号")},
        )
        self.assertFalse(frappe.db.exists("Designation", self.rows[1][0][0]))

    def test_save_failure_rolls_back_and_redacts(self):
        with patch(
            "enterprise_base.permissions.Employee.save",
            side_effect=lambda *a, **kw: frappe.throw("PRIVATE_MARKER"),
        ):
            result = self.run_import()
        self.assertTrue(result["errors"])
        self.assertNotIn("PRIVATE_MARKER", json.dumps(result))
        self.assertNotIn("PRIVATE_MARKER", json.dumps(frappe.local.message_log))
        self.assertFalse(frappe.db.exists("Designation", self.rows[1][0][0]))
        self.assertFalse(frappe.db.exists("Department", {"department_name": self.rows[0][1][0]}))

    def test_other_employee_userid_is_rejected(self):
        other = frappe.copy_doc(self.emp)
        other.employee_number, other.wecom_userid = self.suffix + "b", self.suffix + "b"
        other.insert()
        self.rows[2][0][2] = other.wecom_userid
        self.assertTrue(self.run_import()["errors"])
        self.emp.reload()
        self.assertEqual(self.emp.wecom_userid, self.suffix)

    def test_unprivileged_and_sensitive_headers(self):
        frappe.set_user("Guest")
        try:
            with self.assertRaises(frappe.PermissionError):
                import_workbook("unread", "unread", True)
        finally:
            frappe.set_user("Administrator")
        for field in ("手机号", "身份证号", "银行卡", "薪资"):
            content = self.content(extra=field)
            rows, errors = read_rows(content)
            self.assertTrue(errors, field)
            self.assertNotIn("PRIVATE_MARKER", json.dumps(errors))
